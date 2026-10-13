"""Generates additional candidate resumes (11 through 20) for comprehensive evaluation."""

from pathlib import Path
import docx
import pymupdf

DATA_DIR = Path("data/sample_resumes")
DATA_DIR.mkdir(parents=True, exist_ok=True)

EXTRA_RESUMES = [
    {
        "filename": "resume_11_junior_ml.pdf",
        "format": "pdf",
        "text": """KAI ZHANG
kai.zhang@example.com | Seattle, WA

SUMMARY
Junior Machine Learning Engineer with 1.5 years experience building tabular and computer vision models.

EXPERIENCE
Junior ML Engineer | VisionAI Labs | Aug 2023 - Present
- Trained image classification models using PyTorch, OpenCV, and scikit-learn.
- Built data preprocessing scripts in Python and Pandas.
- Wrote API endpoints using FastAPI and Docker.

EDUCATION
Bachelor of Science in Data Science | University of Washington | 2019 - 2023

SKILLS
Python, PyTorch, Scikit-learn, OpenCV, Pandas, NumPy, FastAPI, Docker, Git, Linux
"""
    },
    {
        "filename": "resume_12_senior_data_scientist.docx",
        "format": "docx",
        "text": """DR. AMANDA FOSTER
amanda.foster@example.com | New York, NY

PROFESSIONAL SUMMARY
Lead Data Scientist with 7 years experience delivering machine learning solutions, experimentation platforms, and causal inference models.

WORK EXPERIENCE
Lead Data Scientist | FinTech Global | Jan 2020 - Present
- Built fraud detection pipelines using Python, LightGBM, XGBoost, and Scikit-learn.
- Designed company-wide A/B testing and experimentation framework.
- Deployed real-time model inference services on AWS with Docker.

Senior Data Analyst | Capital Insight | Jun 2017 - Dec 2019
- Built financial forecasting models in Python and R.
- Developed interactive Tableau and Streamlit dashboards for executives.

EDUCATION
Master of Science in Statistics | Columbia University | 2015 - 2017
Bachelor of Arts in Mathematics | Cornell University | 2011 - 2015

SKILLS
Python, R, SQL, Scikit-learn, LightGBM, XGBoost, Pandas, A/B Testing, AWS, Docker, Tableau, Streamlit
"""
    },
    {
        "filename": "resume_13_senior_react_frontend.txt",
        "format": "txt",
        "text": """LEO DUPONT
leo.dupont@example.com | San Francisco, CA

SUMMARY
Senior Frontend Engineer with 6 years experience architecting large-scale React and Next.js web applications.

EXPERIENCE
Senior Frontend Developer | SaaS Flow | Jan 2021 - Present
- Led frontend architecture migration to Next.js, React, TypeScript, and Redux Toolkit.
- Built design system components using Tailwind CSS and Storybook.
- Improved Core Web Vitals score by 45%.

Frontend Developer | WebCraft | Jun 2018 - Dec 2020
- Developed web interfaces using JavaScript, React, HTML5, CSS3, and REST API integrations.

EDUCATION
Bachelor of Science in Computer Science | San Jose State University | 2014 - 2018

SKILLS
React, Next.js, TypeScript, JavaScript, Redux, Tailwind CSS, HTML5, CSS3, Webpack, Git, JIRA
"""
    },
    {
        "filename": "resume_14_cloud_security_devops.pdf",
        "format": "pdf",
        "text": """ZACHARY MORRIS
zachary.morris@example.com | Dallas, TX

SUMMARY
DevSecOps and Cloud Security Engineer with 5 years experience securing AWS infrastructure and Kubernetes clusters.

EXPERIENCE
Senior Cloud Security Engineer | SecureCloud Inc | Mar 2021 - Present
- Implemented Infrastructure as Code using Terraform and Ansible with automated security linting.
- Managed EKS Kubernetes cluster security, RBAC policies, and container scanning.
- Automated CI/CD security pipelines using GitLab CI and Docker.

DevOps Engineer | CloudOps Solutions | Jun 2019 - Feb 2021
- Maintained AWS cloud infrastructure, VPC networking, and Terraform templates.

EDUCATION
Bachelor of Science in Cybersecurity | University of Texas at Dallas | 2015 - 2019

SKILLS
AWS, Kubernetes, Docker, Terraform, Ansible, CI/CD, GitLab CI, Linux, Bash, Python, Prometheus
"""
    },
    {
        "filename": "resume_15_senior_qa_lead.docx",
        "format": "docx",
        "text": """NATHAN DRAKE
nathan.drake@example.com | Philadelphia, PA

SUMMARY
Senior QA Automation Lead with 8 years experience designing enterprise test frameworks and automated regression suites.

EXPERIENCE
Lead QA Engineer | HealthTech Digital | Jan 2020 - Present
- Architected end-to-end automation test framework using Python, Selenium, Pytest, and Postman.
- Integrated automated testing suites into CI/CD deployment pipelines.
- Mentored a team of 5 QA engineers.

Senior QA Analyst | FinServ | May 2016 - Dec 2019
- Created automated test scripts in Python and Selenium; logged issues in JIRA.

EDUCATION
Bachelor of Science in Information Systems | Temple University | 2012 - 2016

SKILLS
Selenium, Pytest, Python, Postman, JIRA, Unit Testing, CI/CD, Git, Agile
"""
    },
    {
        "filename": "resume_16_nlp_engineer.pdf",
        "format": "pdf",
        "text": """SAMIRA KHAN
samira.khan@example.com | San Francisco, CA

SUMMARY
NLP & Machine Learning Engineer with 4 years experience fine-tuning large language models and developing vector search systems.

EXPERIENCE
Machine Learning Engineer (NLP) | LanguageTech | Jan 2022 - Present
- Fine-tuned BERT and RoBERTa models using PyTorch and Hugging Face Transformers.
- Built semantic search and RAG retrieval pipelines using Sentence-Transformers and FAISS.
- Served real-time inference microservices using FastAPI and Docker on AWS.

Data Scientist | TextMetrics | Jul 2020 - Dec 2021
- Extracted entities and relationships from legal documents using spaCy and Python.
- Built text classification pipelines with scikit-learn and NLTK.

EDUCATION
Master of Science in Artificial Intelligence | Carnegie Mellon University | 2018 - 2020
Bachelor of Science in Computer Science | UC San Diego | 2014 - 2018

SKILLS
Python, PyTorch, Hugging Face, Transformers, Sentence-Transformers, BERT, LLM, RAG, FAISS, spaCy, NLTK, Scikit-learn, FastAPI, Docker, AWS, Git
"""
    },
    {
        "filename": "resume_17_junior_backend_python.txt",
        "format": "txt",
        "text": """MIA ROBERTS
mia.roberts@example.com | Denver, CO

SUMMARY
Junior Backend Software Developer with 2 years experience building REST APIs with Python, Django, and PostgreSQL.

EXPERIENCE
Backend Developer | Peak Web Services | Jan 2023 - Present
- Developed REST API endpoints using Python, Django, and PostgreSQL.
- Implemented background tasks using Celery and Redis.
- Wrote unit tests using Pytest with 85% coverage.

Junior Software Intern | CodeCraft | Jun 2022 - Dec 2022
- Maintained database queries and API documentation with Postman.

EDUCATION
Bachelor of Science in Computer Science | University of Colorado | 2018 - 2022

SKILLS
Python, Django, FastAPI, PostgreSQL, Redis, Celery, REST API, Pytest, Docker, Git, Postman
"""
    },
    {
        "filename": "resume_18_bi_reporting_analyst.docx",
        "format": "docx",
        "text": """BENJAMIN SCOTT
benjamin.scott@example.com | Phoenix, AZ

SUMMARY
Senior Business Intelligence and Data Analyst with 5 years experience in SQL data modeling, ETL, and executive reporting.

EXPERIENCE
Senior BI Analyst | Desert Retail | Jan 2021 - Present
- Built data models and ETL workflows in Snowflake and BigQuery using SQL.
- Designed executive dashboards in Tableau and Power BI.
- Conducted customer segmentation analysis with Python and Pandas.

BI Developer | Metro Supply | Jun 2019 - Dec 2020
- Automated daily financial KPI reporting using Excel, SQL, and Power BI.

EDUCATION
Bachelor of Science in Business Analytics | Arizona State University | 2015 - 2019

SKILLS
SQL, Tableau, Power BI, Snowflake, BigQuery, ETL, Python, Pandas, Excel, Data Modeling
"""
    },
    {
        "filename": "resume_19_intern_software.txt",
        "format": "txt",
        "text": """ETHAN HUNT
ethan.hunt@example.com | Los Angeles, CA

SUMMARY
Undergraduate Software Engineering Intern with coursework in computer science and basic Python programming.

EXPERIENCE
Software Engineering Intern | College Labs | Jun 2024 - Dec 2024
- Assisted with writing unit tests in Python.
- Built a simple internal dashboard using HTML and JavaScript.

EDUCATION
High School Diploma | Beverly Hills High | 2019 - 2023
Currently enrolled in B.S. Computer Science | UCLA (Expected 2027)

SKILLS
Python, JavaScript, HTML5, Git
"""
    },
    {
        "filename": "resume_20_staff_infrastructure_architect.pdf",
        "format": "pdf",
        "text": """VICTORIA STERLING
victoria.sterling@example.com | Seattle, WA

EXECUTIVE SUMMARY
Staff Cloud Infrastructure Architect with 12 years experience designing enterprise Kubernetes platforms, hybrid cloud systems, and IaC pipelines.

EXPERIENCE
Staff Infrastructure Architect | Enterprise Cloud Tech | Jan 2018 - Present
- Designed multi-region Kubernetes (EKS) infrastructure spanning 120+ microservices on AWS.
- Led enterprise infrastructure as code using Terraform, Helm, and Ansible.
- Built unified CI/CD pipelines with GitLab CI and GitHub Actions.
- Mentored 15 DevOps and platform engineers.

Senior Systems Architect | CloudWave | Jan 2013 - Dec 2017
- Managed Linux server clusters and VMware infrastructure.

EDUCATION
Master of Science in Distributed Systems | University of Washington | 2010 - 2012
Bachelor of Science in Computer Science | Purdue University | 2006 - 2010

SKILLS
AWS, Kubernetes, Terraform, Ansible, Docker, Helm, CI/CD, GitLab CI, GitHub Actions, Linux, Prometheus, Grafana, Microservices, Python, Bash
"""
    }
]

def main():
    print(f"Generating resumes 11 through 20 into {DATA_DIR}...")
    for item in EXTRA_RESUMES:
        target_path = DATA_DIR / item["filename"]
        fmt = item["format"]
        text = item["text"].strip()

        if fmt == "txt":
            with open(target_path, "w", encoding="utf-8") as f:
                f.write(text)
        elif fmt == "docx":
            doc = docx.Document()
            for line in text.split("\n"):
                if line.isupper() and len(line) < 40 and not ("|" in line or "@" in line):
                    doc.add_heading(line, level=2)
                else:
                    doc.add_paragraph(line)
            doc.save(str(target_path))
        elif fmt == "pdf":
            doc = pymupdf.open()
            page = doc.new_page(width=612, height=792)
            y = 50
            for line in text.split("\n"):
                if y > 750:
                    page = doc.new_page(width=612, height=792)
                    y = 50
                is_hdr = line.isupper() and len(line) < 40 and not ("|" in line or "@" in line)
                font_sz = 13 if is_hdr else 10
                page.insert_text((50, y), line, fontsize=font_sz)
                y += 18 if is_hdr else 14
            doc.save(str(target_path))
        print(f"Generated: {target_path.name}")

if __name__ == "__main__":
    main()
