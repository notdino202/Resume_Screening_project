"""Document extraction module with PDF (pdfplumber/PyMuPDF), DOCX, TXT, and OCR fallback."""

from __future__ import annotations
from dataclasses import dataclass, field
import io
import os
from pathlib import Path
from typing import Optional

import docx
import pdfplumber
import pymupdf

try:
    import pytesseract
    from PIL import Image
    _HAS_PYTESSERACT = True
except ImportError:
    _HAS_PYTESSERACT = False

from .cleaner import CleanedTextResult, TextCleaner
from .hasher import compute_bytes_hash, compute_file_hash


class ExtractionError(Exception):
    """Raised when document ingestion or parsing fails."""
    pass


@dataclass
class ExtractionResult:
    """Encapsulates raw extraction, cleaning, and diagnostic metadata."""
    file_path: Optional[str]
    file_name: str
    file_hash: str
    file_type: str
    raw_text: str
    cleaned_result: CleanedTextResult
    method_used: str
    ocr_attempted: bool = False
    ocr_error: Optional[str] = None
    extraction_error: Optional[str] = None
    is_success: bool = True
    advisory_flags: list[str] = field(default_factory=list)

    @property
    def clean_text(self) -> str:
        return self.cleaned_result.cleaned_text


class DocumentExtractor:
    """Robust document extractor supporting PDF, DOCX, and TXT with fallbacks."""

    def __init__(self, cleaner: Optional[TextCleaner] = None, min_text_for_ocr: int = 50):
        self.cleaner = cleaner or TextCleaner()
        self.min_text_for_ocr = min_text_for_ocr

    def extract(self, file_path_or_bytes: str | Path | bytes, file_name: Optional[str] = None) -> ExtractionResult:
        """Extracts text from a file path or in-memory bytes.

        Args:
            file_path_or_bytes: Path to file or bytes payload.
            file_name: Optional file name when passing raw bytes.

        Returns:
            ExtractionResult containing cleaned text, hash, method used, and quality flags.

        Raises:
            ExtractionError: If file is empty, corrupted, or unsupported.
        """
        if isinstance(file_path_or_bytes, (str, Path)):
            file_path = Path(file_path_or_bytes)
            if not file_path.exists():
                raise ExtractionError(f"File not found: {file_path}")
            if file_path.stat().st_size == 0:
                raise ExtractionError(f"Empty file (0 bytes): {file_path.name}")
            with open(file_path, "rb") as f:
                content_bytes = f.read()
            target_name = file_name or file_path.name
            path_str = str(file_path)
        elif isinstance(file_path_or_bytes, bytes):
            content_bytes = file_path_or_bytes
            if len(content_bytes) == 0:
                raise ExtractionError(f"Empty document bytes provided ({file_name or 'unnamed'})")
            target_name = file_name or "uploaded_document"
            path_str = None
        else:
            raise ExtractionError(f"Unsupported input type: {type(file_path_or_bytes)}")

        file_hash = compute_bytes_hash(content_bytes)
        ext = Path(target_name).suffix.lower()

        if ext == ".txt":
            return self._extract_txt(content_bytes, target_name, path_str, file_hash)
        elif ext in (".docx", ".doc"):
            if ext == ".doc":
                raise ExtractionError("Legacy binary .doc format not supported. Please convert to .docx or .pdf.")
            return self._extract_docx(content_bytes, target_name, path_str, file_hash)
        elif ext == ".pdf":
            return self._extract_pdf(content_bytes, target_name, path_str, file_hash)
        else:
            raise ExtractionError(f"Unsupported document format '{ext}'. Supported formats: .pdf, .docx, .txt")

    def _extract_txt(self, content_bytes: bytes, file_name: str, file_path: Optional[str], file_hash: str) -> ExtractionResult:
        """Extracts and decodes plain text with encoding fallbacks."""
        text = ""
        for encoding in ("utf-8", "latin-1", "cp1252"):
            try:
                text = content_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        if not text and content_bytes:
            raise ExtractionError(f"Unable to decode text file: {file_name}")

        cleaned = self.cleaner.clean(text)
        flags = list(cleaned.quality_flags)

        return ExtractionResult(
            file_path=file_path,
            file_name=file_name,
            file_hash=file_hash,
            file_type="TXT",
            raw_text=text,
            cleaned_result=cleaned,
            method_used="text_plain",
            advisory_flags=flags,
        )

    def _extract_docx(self, content_bytes: bytes, file_name: str, file_path: Optional[str], file_hash: str) -> ExtractionResult:
        """Extracts text from DOCX paragraphs and tables."""
        try:
            doc = docx.Document(io.BytesIO(content_bytes))
        except Exception as e:
            raise ExtractionError(f"Corrupted or invalid DOCX file '{file_name}': {str(e)}") from e

        text_pieces = []
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_pieces.append(paragraph.text)

        # Also extract table text to avoid missing structured resume sections
        for table in doc.tables:
            for row in table.rows:
                row_texts = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_texts:
                    text_pieces.append(" | ".join(row_texts))

        raw_text = "\n\n".join(text_pieces)
        cleaned = self.cleaner.clean(raw_text)
        flags = list(cleaned.quality_flags)

        return ExtractionResult(
            file_path=file_path,
            file_name=file_name,
            file_hash=file_hash,
            file_type="DOCX",
            raw_text=raw_text,
            cleaned_result=cleaned,
            method_used="python-docx",
            advisory_flags=flags,
        )

    def _extract_pdf(self, content_bytes: bytes, file_name: str, file_path: Optional[str], file_hash: str) -> ExtractionResult:
        """Extracts text from PDF using pdfplumber, falling back to PyMuPDF and OCR."""
        raw_text = ""
        method_used = "pdfplumber"
        ocr_attempted = False
        ocr_error = None

        # 1. Primary: pdfplumber
        try:
            with pdfplumber.open(io.BytesIO(content_bytes)) as pdf:
                pages_text = []
                for page in pdf.pages:
                    pt = page.extract_text(layout=True) or page.extract_text() or ""
                    if pt.strip():
                        pages_text.append(pt)
                raw_text = "\n\n".join(pages_text).strip()
        except Exception as e:
            # Fallback to PyMuPDF if pdfplumber fails
            method_used = "pymupdf_fallback"
            try:
                doc = pymupdf.open(stream=content_bytes, filetype="pdf")
                pages_text = []
                for page in doc:
                    pt = page.get_text() or ""
                    if pt.strip():
                        pages_text.append(pt)
                raw_text = "\n\n".join(pages_text).strip()
            except Exception as e2:
                raise ExtractionError(f"Corrupted or invalid PDF file '{file_name}': {str(e2)}") from e2

        # If pdfplumber gave empty text, try PyMuPDF as well before OCR
        if not raw_text.strip() and method_used == "pdfplumber":
            try:
                doc = pymupdf.open(stream=content_bytes, filetype="pdf")
                pages_text = []
                for page in doc:
                    pt = page.get_text() or ""
                    if pt.strip():
                        pages_text.append(pt)
                if pages_text:
                    raw_text = "\n\n".join(pages_text).strip()
                    method_used = "pymupdf_fallback"
            except Exception:
                pass

        # 2. OCR Fallback if extracted text is near-empty (likely scanned document)
        if len(raw_text.strip()) < self.min_text_for_ocr:
            ocr_attempted = True
            ocr_text = self._attempt_ocr(content_bytes)
            if ocr_text:
                raw_text = ocr_text
                method_used = "ocr_fallback"
            else:
                ocr_error = "OCR unavailable or produced no text"

        cleaned = self.cleaner.clean(raw_text)
        flags = list(cleaned.quality_flags)

        if ocr_attempted and method_used != "ocr_fallback":
            flags.append("scanned-resume-unextracted")
            if "low-quality extraction" not in flags:
                flags.append("low-quality extraction")

        return ExtractionResult(
            file_path=file_path,
            file_name=file_name,
            file_hash=file_hash,
            file_type="PDF",
            raw_text=raw_text,
            cleaned_result=cleaned,
            method_used=method_used,
            ocr_attempted=ocr_attempted,
            ocr_error=ocr_error,
            advisory_flags=flags,
        )

    def _attempt_ocr(self, content_bytes: bytes) -> Optional[str]:
        """Attempts OCR on rendered PDF pages if Tesseract is available."""
        if not _HAS_PYTESSERACT:
            return None

        try:
            doc = pymupdf.open(stream=content_bytes, filetype="pdf")
            extracted_pages = []
            for page in doc:
                pix = page.get_pixmap(dpi=150)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                text = pytesseract.image_to_string(img)
                if text.strip():
                    extracted_pages.append(text)
            if extracted_pages:
                return "\n\n".join(extracted_pages).strip()
        except Exception:
            # Tesseract binary not installed or failed
            return None
        return None
